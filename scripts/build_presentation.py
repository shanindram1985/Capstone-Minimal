"""Build the project presentation. Run: .\\.venv\\Scripts\\python.exe scripts\\build_presentation.py"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0E, 0x1C, 0x36)
TEAL = RGBColor(0x1F, 0x8A, 0x70)
GOLD = RGBColor(0xE3, 0xB2, 0x3C)
INK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5D, 0x6B, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CLOUD = RGBColor(0xF4, 0xF7, 0xFB)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE2, 0xE8, 0xF0)

W, H = Inches(13.333), Inches(7.5)


def _set_run(run, text, size, bold, color, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def textbox(slide, text, x, y, w, h, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    _set_run(run, text, size, bold, color, font)
    return box


def rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def footer(slide, page, total=7):
    textbox(slide, "Mobile Script Generator  ·  Wireframe", Inches(0.55), Inches(7.12), Inches(8), Inches(0.28), 11, False, MUTED)
    textbox(slide, f"{page}  /  {total}", Inches(11.4), Inches(7.12), Inches(1.4), Inches(0.28), 11, False, MUTED, PP_ALIGN.RIGHT)


def header_bar(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()


def title_block(slide, kicker, title):
    textbox(slide, kicker.upper(), Inches(0.55), Inches(0.32), Inches(10), Inches(0.32), 13, True, TEAL)
    textbox(slide, title, Inches(0.55), Inches(0.62), Inches(12), Inches(0.55), 32, True, NAVY, font="Calibri")


def bullet_card(slide, x, y, w, h, heading, lines, accent=TEAL):
    rect(slide, x, y, w, h, CARD, LINE)
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.08), h)
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent
    accent_bar.line.fill.background()
    textbox(slide, heading, x + Inches(0.28), y + Inches(0.16), w - Inches(0.4), Inches(0.36), 16, True, NAVY)
    box = slide.shapes.add_textbox(x + Inches(0.28), y + Inches(0.55), w - Inches(0.45), h - Inches(0.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run(run, line, 14, False, INK)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), H)
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    textbox(s, "RETAIL MOBILE QUALITY", Inches(0.7), Inches(1.55), Inches(11), Inches(0.35), 14, True, GOLD)
    textbox(s, "Mobile Script Generator", Inches(0.7), Inches(2.05), Inches(11.5), Inches(0.8), 44, True, WHITE)
    textbox(s, "Wireframe", Inches(0.7), Inches(2.85), Inches(11), Inches(0.7), 40, True, TEAL)
    textbox(
        s,
        "A wireframe of a shop screen becomes a full shopper test:\nwhat should work, what should be refused, and what sits at the edges.",
        Inches(0.7), Inches(3.85), Inches(10.5), Inches(0.9), 20, False, RGBColor(0xD5, 0xDE, 0xEA),
    )
    textbox(s, "Shanmuganathan    ·    September 2026", Inches(0.7), Inches(6.4), Inches(8), Inches(0.35), 14, False, RGBColor(0xA8, 0xB4, 0xC4))

    # 2 Agents
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "The agents", "Five roles. Each one does one job for the shopper journey.")
    agents = [
        ("Vision", "Looks at the wireframe and names the screen, its purpose, and every control a shopper can see."),
        ("Scenario designer", "Turns that screen into shopping situations: success, refusal, and unusual input."),
        ("Script writer", "Writes one check for each situation, ready to run on the store app."),
        ("Locator finder", "On the live phone, finds the button or field by what it says, not by a fixed map."),
        ("Reporter", "Tells the team what passed, what failed, how long it took, and shows the screen when something breaks."),
    ]
    for i, (name, detail) in enumerate(agents):
        y = Inches(1.45 + i * 1.05)
        rect(s, Inches(0.5), y, Inches(12.3), Inches(0.95), CARD, LINE)
        pill = rect(s, Inches(0.7), y + Inches(0.24), Inches(2.5), Inches(0.46), TEAL if i % 2 == 0 else NAVY)
        textbox(s, name, Inches(0.7), y + Inches(0.3), Inches(2.5), Inches(0.36), 14, True, WHITE, PP_ALIGN.CENTER)
        textbox(s, detail, Inches(3.45), y + Inches(0.22), Inches(9.1), Inches(0.55), 16, False, INK)
    footer(s, 2)

    # 3 Pipeline
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "The pipeline", "From a picture of the shop to a result the business can read")
    steps = [
        ("1", "See the screen", "A wireframe of login, catalog, product, cart, or checkout"),
        ("2", "Understand it", "What the screen is for, and which controls matter"),
        ("3", "Design checks", "A mix of good, bad, and edge shopper behavior"),
        ("4", "Run the shop", "Each check is played on the live demo store"),
        ("5", "Show the outcome", "A clear report, not a raw log"),
    ]
    for i, (num, name, detail) in enumerate(steps):
        x = Inches(0.4 + i * 2.58)
        rect(s, x, Inches(1.7), Inches(2.42), Inches(3.55), CARD, LINE)
        bubble = s.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.88), Inches(1.95), Inches(0.62), Inches(0.62))
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = TEAL if i < 4 else GOLD
        bubble.line.fill.background()
        textbox(s, num, x + Inches(0.88), Inches(2.06), Inches(0.62), Inches(0.42), 18, True, WHITE, PP_ALIGN.CENTER)
        textbox(s, name, x + Inches(0.12), Inches(2.75), Inches(2.18), Inches(0.7), 16, True, NAVY, PP_ALIGN.CENTER)
        textbox(s, detail, x + Inches(0.16), Inches(3.5), Inches(2.1), Inches(1.4), 14, False, MUTED, PP_ALIGN.CENTER)
    textbox(
        s,
        "Review the checks first, with no phone. Run them when the store app and a device are ready.",
        Inches(0.55), Inches(5.55), Inches(12.2), Inches(0.5), 16, False, INK,
    )
    footer(s, 3)

    # 4 Scenario combinations
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "Scenario mix", "Every screen is tested as a shopper would actually use it")
    textbox(
        s,
        "The script writer does not stop at “it works.” Each wireframe gets three kinds of check, using only what is on that screen.",
        Inches(0.55), Inches(1.38), Inches(12.2), Inches(0.45), 16, False, MUTED,
    )
    kinds = [
        ("Positive", "The happy shopper", "Right account, a real product search, quantity 1, a valid offer, a complete checkout."),
        ("Negative", "The blocked shopper", "Locked account, wrong password, empty search, a fake offer code, a missing card."),
        ("Edge", "The unusual shopper", "Blank spaces, very long text, quantity 0 and 99, odd characters in a field."),
    ]
    for i, (kind, subtitle, detail) in enumerate(kinds):
        x = Inches(0.45 + i * 4.25)
        rect(s, x, Inches(2.0), Inches(4.05), Inches(2.55), CARD, LINE)
        textbox(s, kind, x + Inches(0.22), Inches(2.15), Inches(3.6), Inches(0.4), 20, True, TEAL if i != 1 else RGBColor(0xC4, 0x45, 0x45))
        textbox(s, subtitle, x + Inches(0.22), Inches(2.58), Inches(3.6), Inches(0.35), 14, True, NAVY)
        textbox(s, detail, x + Inches(0.22), Inches(3.05), Inches(3.6), Inches(1.25), 14, False, INK)
    rows = [
        ("Sign in", "Known shoppers and a visual user", "Locked out, bad password, blanks", "Spaces and oversized text"),
        ("Browse", "Find “backpack”, scroll the catalog", "Empty search, no such product", "A very long or odd search"),
        ("Buy", "Add one item, valid offer, pay", "Zero quantity, fake offer, missing card", "Quantity 99 and odd address text"),
    ]
    for i, (area, pos, neg, edge) in enumerate(rows):
        y = Inches(4.75 + i * 0.7)
        textbox(s, area, Inches(0.5), y, Inches(1.4), Inches(0.4), 14, True, NAVY)
        textbox(s, pos, Inches(2.0), y, Inches(3.5), Inches(0.55), 13, False, INK)
        textbox(s, neg, Inches(5.6), y, Inches(3.6), Inches(0.55), 13, False, INK)
        textbox(s, edge, Inches(9.3), y, Inches(3.5), Inches(0.55), 13, False, INK)
    footer(s, 4)

    # 5 How locators are chosen
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "Finding the control", "The test looks at the live shop, then chooses what to tap")
    steps = [
        ("1", "Open the screen", "The shop app is on the phone. A pop-up from the phone itself is closed first so it cannot hide the store."),
        ("2", "Read what is visible", "The finder looks at the words on the screen: Menu, Log in, search, price, Add to cart."),
        ("3", "Match the wireframe", "It pairs each step with the control that says the same thing the wireframe named."),
        ("4", "Prefer the real store", "If the demo shop’s own menu, login, or cart control is on screen, that one is used."),
        ("5", "Remember the choice", "The match is kept for that check and written down, so the report can show what was used."),
    ]
    for i, (num, name, detail) in enumerate(steps):
        y = Inches(1.45 + i * 1.05)
        rect(s, Inches(0.5), y, Inches(12.3), Inches(0.95), CARD, LINE)
        bubble = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), y + Inches(0.16), Inches(0.62), Inches(0.62))
        bubble.fill.solid()
        bubble.fill.fore_color.rgb = TEAL
        bubble.line.fill.background()
        textbox(s, num, Inches(0.7), y + Inches(0.28), Inches(0.62), Inches(0.4), 16, True, WHITE, PP_ALIGN.CENTER)
        textbox(s, name, Inches(1.55), y + Inches(0.12), Inches(3.3), Inches(0.7), 16, True, NAVY)
        textbox(s, detail, Inches(4.9), y + Inches(0.16), Inches(7.6), Inches(0.65), 15, False, INK)
    footer(s, 5)

    # 6 Reporter
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "The report", "A result a product owner can read without opening a log")
    cards = [
        ("The headline", "Passed, or needs attention, with the share of checks that succeeded."),
        ("The numbers", "How many ran, passed, failed, were skipped, and how long the run took."),
        ("The pictures", "A ring of outcomes, and a bar showing which check took the longest."),
        ("The list", "Each check in plain words, tagged Positive, Negative, or Edge."),
        ("The evidence", "If a check fails, the reason opens under the row, with a photo of the phone."),
        ("The audience", "QA, product, and demo reviewers see the same page. It opens when the run finishes."),
    ]
    for i, (head, body) in enumerate(cards):
        col = i % 3
        row = i // 3
        x = Inches(0.45 + col * 4.25)
        y = Inches(1.55 + row * 2.55)
        rect(s, x, y, Inches(4.05), Inches(2.35), CARD, LINE)
        textbox(s, head, x + Inches(0.22), y + Inches(0.22), Inches(3.6), Inches(0.45), 18, True, NAVY)
        textbox(s, body, x + Inches(0.22), y + Inches(0.8), Inches(3.6), Inches(1.3), 15, False, INK)
    footer(s, 6)

    # 7 Business outcome
    s = prs.slides.add_slide(blank)
    header_bar(s)
    title_block(s, "Why it matters", "Coverage follows the wireframe, not a single happy path")
    bullet_card(
        s, Inches(0.5), Inches(1.55), Inches(6.15), Inches(5.15),
        "For the business",
        [
            "A new shop screen can be tested from its wireframe.",
            "Login, browse, product, cart, and checkout each get their own mix.",
            "Risk shows up as a refused login, a bad offer, or a broken checkout, not as one green login.",
            "The report is ready to show in a review.",
        ],
    )
    bullet_card(
        s, Inches(6.9), Inches(1.55), Inches(5.9), Inches(5.15),
        "What is in place",
        [
            "Vision, scenario design, script writing, live control finding, and reporting.",
            "Retail examples: real shopper, locked account, backpack search, quantity, offer code, payment fields.",
            "Demo store: Sauce Labs My Demo App.",
            "Same story for any later wireframe of the same shop.",
        ],
        GOLD,
    )
    footer(s, 7)

    out = Path(__file__).resolve().parents[1] / "Mobile-Script-Generator-Wireframe.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    build()
